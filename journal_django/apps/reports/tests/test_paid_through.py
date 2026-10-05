"""
Тесты отчёта «Оплачено до» (apps.reports.builders.paid_through).

Лежат в apps/reports/tests намеренно: пакет поднимает свежую test_journal_test
(см. conftest.py рядом), а отчёт читает ВСЕХ учеников базы — на общей
persistent journal_test состав был бы недетерминирован.
"""
from __future__ import annotations

import datetime
import io
from decimal import Decimal

import openpyxl
import pytest

from apps.directions.models import Direction
from apps.groups.models import Group
from apps.lessons.models import Lesson, LessonAttendance
from apps.memberships.models import GroupMembership
from apps.payments.models import Payment
from apps.reports import services
from apps.reports.builders.paid_through import build, collect, paid_until
from apps.reports.models import ReportType
from apps.students.models import Student
from apps.teachers.models import Teacher

pytestmark = pytest.mark.django_db

CREATED_AT = datetime.datetime(2026, 1, 1, tzinfo=datetime.UTC)
TODAY = datetime.date(2026, 10, 5)


# --- paid_until: чистая арифметика «4 урока = 1 месяц» ----------------------

@pytest.mark.parametrize('balance, months, until', [
    # Ровно абонемент — ровно месяц вперёд.
    (Decimal('4'), Decimal('1'), datetime.date(2026, 11, 5)),
    # Целые месяцы — календарные, не по 30 дней.
    (Decimal('12'), Decimal('3'), datetime.date(2027, 1, 5)),
    # Остаток уроков сверх целых месяцев — по неделе на урок.
    (Decimal('6'), Decimal('1.5'), datetime.date(2026, 11, 19)),
    # Пол-урока (45 мин) — пол-недели, округление вверх до дня.
    (Decimal('0.5'), Decimal('0.13'), datetime.date(2026, 10, 9)),
])
def test_paid_until_four_lessons_per_month(balance, months, until):
    assert paid_until(balance, TODAY) == (months, until)


@pytest.mark.parametrize('balance', [Decimal('0'), Decimal('-3'), Decimal('-0.5')])
def test_paid_until_without_paid_lessons_has_no_date(balance):
    """Нет оплаченных уроков — учёба не оплачена: 0 месяцев, даты нет."""
    assert paid_until(balance, TODAY) == (Decimal('0'), None)


def test_paid_until_month_end_clamps_to_last_day():
    """31 января + 1 месяц — последний день февраля, а не 3 марта."""
    assert paid_until(Decimal('4'), datetime.date(2026, 1, 31)) == (
        Decimal('1'), datetime.date(2026, 2, 28))


# --- collect: строки отчёта ------------------------------------------------

@pytest.fixture
def data():
    class F:
        def __init__(self):
            self.teacher = Teacher.objects.create(name='__pt_teacher__', created_at=CREATED_AT)

        def direction(self, name: str) -> Direction:
            return Direction.objects.create(name=name, total_lessons=8, active=True)

        def group(self, name: str, direction: Direction, active: bool = True) -> Group:
            return Group.objects.create(
                name=name, direction=direction, teacher=self.teacher, active=active,
                is_individual=False, lesson_duration_minutes=90, created_at=CREATED_AT,
            )

        def student(self, full_name: str, platform_id: str | None = None) -> Student:
            return Student.objects.create(
                full_name=full_name, platform_id=platform_id, created_at=CREATED_AT)

        def member(self, student: Student, group: Group, active: bool = True):
            return GroupMembership.objects.create(student=student, group=group, active=active)

        def pay(self, student: Student, direction: Direction, lessons: int):
            return Payment.objects.create(
                student=student, direction=direction, subscriptions_count=lessons // 4,
                lessons_count=lessons, unit_price=Decimal('4000.00'),
                total_amount=Decimal('4000.00') * (lessons // 4),
                paid_at='2026-09-01', created_at=CREATED_AT,
            )

        def attend(self, student: Student, group: Group, date: str, number: int):
            lesson = Lesson.objects.create(
                group=group, teacher=self.teacher, lesson_date=date, lesson_number=number,
                lesson_duration_minutes=90, lesson_type='regular',
                submitted_at=CREATED_AT, submitted_by_token='__pt_test__',
            )
            LessonAttendance.objects.create(lesson=lesson, student=student, present=True)

    return F()


def _rows():
    return {r.full_name: r for r in collect(TODAY) if r.full_name.startswith('__pt_')}


def test_row_has_ids_course_balance_and_paid_until(data):
    d = data.direction('__pt_Python__')
    g = data.group('__pt_g1__', d)
    s = data.student('__pt_s1__', platform_id='P-100')
    data.member(s, g)
    data.pay(s, d, 8)
    data.attend(s, g, '2026-09-07', 1)
    data.attend(s, g, '2026-09-14', 2)

    row = _rows()['__pt_s1__']

    assert (row.student_id, row.platform_id, row.courses) == (s.id, 'P-100', '__pt_Python__')
    assert (row.balance, row.months, row.paid_until) == (
        Decimal('6'), Decimal('1.5'), datetime.date(2026, 11, 19))


def test_two_active_groups_give_one_row_with_both_courses(data):
    s = data.student('__pt_s2__')
    data.member(s, data.group('__pt_g2a__', data.direction('__pt_Roblox__')))
    data.member(s, data.group('__pt_g2b__', data.direction('__pt_Blender__')))

    assert _rows()['__pt_s2__'].courses == '__pt_Blender__, __pt_Roblox__'


def test_inactive_membership_or_inactive_group_is_not_current_course(data):
    """Курс сейчас — только активное членство в АКТИВНОЙ группе: группы,
    восстановленные из истории, неактивны, хотя членство в них осталось."""
    d = data.direction('__pt_Minecraft__')
    s = data.student('__pt_s3__')
    data.member(s, data.group('__pt_g3a__', d), active=False)
    data.member(s, data.group('__pt_g3b__', d, active=False))

    assert _rows()['__pt_s3__'].courses == ''


def test_student_without_groups_and_payments_is_in_report(data):
    """Отчёт по всем ученикам базы: без групп и оплат — строка с нулями."""
    data.student('__pt_s4__')

    row = _rows()['__pt_s4__']

    assert (row.courses, row.balance, row.months, row.paid_until) == (
        '', Decimal('0'), Decimal('0'), None)


def test_debt_has_negative_balance_and_no_date(data):
    d = data.direction('__pt_Web__')
    g = data.group('__pt_g5__', d)
    s = data.student('__pt_s5__')
    data.member(s, g)
    data.attend(s, g, '2026-09-07', 1)

    row = _rows()['__pt_s5__']

    assert (row.balance, row.months, row.paid_until) == (Decimal('-1'), Decimal('0'), None)


def test_rows_sorted_by_student_id(data):
    first = data.student('__pt_s6a__')
    second = data.student('__pt_s6b__')

    ids = [r.student_id for r in collect(TODAY) if r.full_name.startswith('__pt_s6')]

    assert ids == [first.id, second.id]


def test_build_writes_sheet_with_headers(data):
    s = data.student('__pt_s7__', platform_id='P-7')
    data.pay(s, data.direction('__pt_Scratch__'), 4)

    content, row_count, filename = build()

    ws = openpyxl.load_workbook(io.BytesIO(content)).active
    assert [c.value for c in ws[1]] == [
        'ID ученика', 'Platform ID', 'Ученик', 'Курс сейчас',
        'Остаток уроков', 'Месяцев обучения', 'Оплачено до',
    ]
    assert row_count == ws.max_row - 1
    assert filename.startswith('paid_through_') and filename.endswith('.xlsx')
    row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[2] == '__pt_s7__')
    assert row[0] == s.id and row[1] == 'P-7' and row[4] == 4 and row[5] == 1
    assert isinstance(row[6], datetime.datetime)


def test_service_dispatches_report_type(data):
    content, _, filename = services.build_report(ReportType.PAID_THROUGH, {})

    assert content[:2] == b'PK' and filename.startswith('paid_through_')
